"""Conferência visual entre cabeçalhos ABRIL (linha 2) e MAIO (linha 3) da folha 05."""
from __future__ import annotations

import shutil
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "02-Referencias" / "folha 05.xlsx"
DST = ROOT / "02-Referencias" / "folha 05_conferencia.xlsx"

RED = PatternFill("solid", fgColor="FF1A1A")
RED_FONT = Font(color="FFFFFF", bold=True)

PASTELS = [
    "FFE8F4FD", "FFF3E8FD", "FFE8FDF0", "FFFFF8E1", "FFFFE8EC",
    "FFE8ECFF", "FFFFF0E8", "FFE8FFF8", "FFF0E8FF", "FFE8FFE8",
    "FFFFE8F8", "FFF8FFE8", "FFE8F8FF", "FFFFF5E8", "FFF5E8FF",
    "FFE8FFF5", "FFFFE8E8", "FFE8E8FF", "FFFFF0F0", "FFF0FFFF",
    "FFFFE0F0", "FFF0E0FF", "FFE0FFF0", "FFFFF0E0", "FFE0F0FF",
    "FFFFE8D6", "FFD6F5FF", "FFD6FFD6", "FFFFD6D6", "FFD6D6FF",
    "FFFFF2CC", "FFCCF2FF", "FFCCFFCC", "FFFFCCCC", "FFCCCCFF",
    "FFE6CCFF", "FFCCE6FF", "FFFFE6CC", "FFCCE6CC", "FFE6E6CC",
    "FFCCFFE6", "FFE6CCF2", "FFF2CCE6", "FFCCE6F2", "FFF2E6CC",
    "FFCCE6E6", "FFE6F2CC", "FFD0E8FF", "FFE8D0FF", "FFFFD0E8",
    "FFD0FFE8", "FFE8FFD0", "FFC8E6FF", "FFFFC8E6", "FFC8FFE6",
    "FFE6FFC8", "FFFFE6C8", "FFC8D6FF", "FFFFC8D6", "FFD6FFC8",
    "FFC8FFD6", "FFFFD6C8", "FFE0C8FF", "FFFFE0C8", "FFC8FFE0",
]

# Pareamento manual ABR(col) -> MAI(col). Conferido campo a campo.
PAIRS: dict[int, int] = {
    1: 1, 2: 2, 3: 3, 4: 4, 5: 5, 6: 6, 7: 7,
    9: 9, 10: 10, 11: 11, 12: 12, 13: 13, 14: 14, 15: 15, 16: 16,
    17: 30,   # AUX.MATERNID <-> AUX MATERNIDADE
    18: 17,   # FERIAS
    19: 47,   # ABONO PEC <-> AB PECUNIARIO
    20: 18, 21: 19, 22: 20, 23: 21,
    25: 22, 26: 23, 27: 24, 28: 25, 29: 26, 30: 27, 31: 28, 32: 29,
    33: 31,   # TT EV TRIB <-> TT EV TRIBUTADOS
    35: 32, 36: 33, 37: 34, 38: 35, 39: 36, 40: 37, 41: 38,
    42: 39, 43: 40, 46: 41, 47: 42, 48: 45, 49: 46,
    50: 48, 51: 49, 52: 43, 53: 50,
    57: 51, 58: 52, 59: 53, 60: 54, 61: 55, 62: 56, 63: 57,
    64: 58, 65: 59, 66: 60, 67: 61, 68: 62, 69: 63,
}

ONLY_ABRIL = {8, 24, 34, 44, 45, 54, 55, 56}
ONLY_MAIO = {44}  # segunda coluna PRODUCAO (col 32 ja pareada com P.PRODUCAO)


def colorize() -> dict:
    shutil.copy2(SRC, DST)
    df = pd.read_excel(SRC, sheet_name=0, header=None)

    abr_row = 2
    mai_row = 3
    abr = df.iloc[1, 1:].tolist()
    mai = df.iloc[2, 1:].tolist()

    wb = load_workbook(DST)
    ws = wb.active

    color_idx = 0
    for ac, mc in sorted(PAIRS.items()):
        fill = PatternFill("solid", fgColor=PASTELS[color_idx % len(PASTELS)])
        color_idx += 1
        ws.cell(abr_row, ac + 1).fill = fill
        ws.cell(mai_row, mc + 1).fill = fill

    for col in ONLY_ABRIL:
        ws.cell(abr_row, col + 1).fill = RED
        ws.cell(abr_row, col + 1).font = RED_FONT

    for col in ONLY_MAIO:
        ws.cell(mai_row, col + 1).fill = RED
        ws.cell(mai_row, col + 1).font = RED_FONT

    # Col 8 MAIO vazia (TOMADOR sumiu)
    if pd.isna(mai[7]):
        ws.cell(mai_row, 8 + 1).fill = RED
        ws.cell(mai_row, 8 + 1).font = RED_FONT

    ws.cell(5, 1, "LEGENDA")
    ws.cell(6, 1, "Mesma cor pastel = mesmo evento nas duas folhas")
    ws.cell(7, 1, "Vermelho forte = evento ausente ou sem par na outra folha")

    wb.save(DST)

    return {"abr": abr, "mai": mai}


def print_report(data: dict) -> None:
    abr = data["abr"]
    mai = data["mai"]

    print("ARQUIVO:", DST)
    print()
    print(f"PARES CONFERIDOS: {len(PAIRS)}")
    for ac in sorted(PAIRS):
        mc = PAIRS[ac]
        print(f"  Col ABR {ac:2d} <-> Col MAI {mc:2d}  |  {abr[ac-1]}  <->  {mai[mc-1]}")
    print()
    print(f"SOMENTE EM ABRIL (vermelho): {len(ONLY_ABRIL)}")
    for col in sorted(ONLY_ABRIL):
        print(f"  Col {col:2d}: {abr[col-1]}")
    print()
    print("SOMENTE EM MAIO (vermelho):")
    print(f"  Col  8: (vazio — TOMADOR existia em abril)")
    for col in sorted(ONLY_MAIO):
        print(f"  Col {col:2d}: {mai[col-1]}")


if __name__ == "__main__":
    print_report(colorize())
