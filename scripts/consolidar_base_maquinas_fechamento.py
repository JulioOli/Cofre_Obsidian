"""
Consolida a base transacional de custos de máquinas (jan/2025 – mar/2026).

Fontes (em ordem de prioridade para rastreio):
  1. Arquivos de fechamento ODBC/SAGI (fechamento2025 + FECHAMENTO_ODBC 2026_01/02/03)
  2. Diesel jan–mar/2026 (relatorio_custoComb_placas_jan-mar_fechamento.xlsx)
  3. Linhas já existentes na base (importação SAGI/Atua anterior)

Saída:
  - 02-Referencias/Custos-Maquinas/relatorio_marcas_maquinas_base.xlsx
    • aba base_maquinas_2025_2026 (com colunas Arquivo_Fonte, Tipo_Fonte)
    • aba manifesto_fontes (resumo por arquivo)

Uso:
  .\\.venv\\Scripts\\python.exe scripts\\consolidar_base_maquinas_fechamento.py
"""

from __future__ import annotations

import re
import sys
from datetime import date
from pathlib import Path

import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT / "scripts") not in sys.path:
    sys.path.insert(0, str(ROOT / "scripts"))

from gerar_analise_contratos_diesel_maquinas import (  # noqa: E402
    carregar_catalogo_maquinas,
    inferir_marca,
    normalizar_placa,
    parse_numero_br,
)

CUSTOS = ROOT / "02-Referencias" / "Custos-Maquinas"
XLSX_BASE = CUSTOS / "relatorio_marcas_maquinas_base.xlsx"
SHEET_BASE = "base_maquinas_2025_2026"
SHEET_MANIFESTO = "manifesto_fontes"

PERIODO_INI = date(2025, 1, 1)
PERIODO_FIM = date(2026, 3, 31)

FECHAMENTO_SOURCES: list[tuple[str, Path]] = [
    ("fechamento2025.xlsx", ROOT / "02-Referencias" / "fechamento2025.xlsx"),
    ("FECHAMENTO_ODBC_2026_01.xlsx", ROOT / "outputs" / "tabelas" / "Fechamento" / "FECHAMENTO_ODBC_2026_01.xlsx"),
    ("FECHAMENTO_ODBC_2026_02.xlsx", ROOT / "outputs" / "tabelas" / "Fechamento" / "FECHAMENTO_ODBC_2026_02.xlsx"),
    ("FECHAMENTO_ODBC_2026_03.xlsx", ROOT / "02-Referencias" / "Fechamento" / "FECHAMENTO_ODBC_2026_03.xlsx"),
]

DIESEL_FONTE = (
    "relatorio_custoComb_placas_jan-mar_fechamento.xlsx",
    CUSTOS / "relatorio_custoComb_placas_jan-mar_fechamento.xlsx",
)

TIPO_FECHAMENTO = "fechamento_odbc"
TIPO_DIESEL = "diesel_abastecimento_q1_2026"
TIPO_SAGI_ATUA = "importacao_sagi_atua"


def extrair_id_maquina(valor) -> str | None:
    s = re.sub(r"\s+", "", str(valor).strip().upper())
    m = re.match(r"^(EHL\d+|EHH\d+)", s)
    return m.group(1) if m else None


def norm_valor_chave(val) -> str:
    parsed = parse_numero_br(val)
    if parsed is not None:
        return f"{parsed:.2f}"
    n = pd.to_numeric(val, errors="coerce")
    if pd.isna(n):
        return str(val).strip()
    return f"{n:.2f}"


def chave_lancamento(df: pd.DataFrame) -> pd.Series:
    d = pd.to_datetime(df["data_nf"], errors="coerce", dayfirst=True)
    return (
        d.dt.strftime("%Y-%m-%d").fillna("")
        + "|"
        + df["cod_conta"].astype(str).str.strip()
        + "|"
        + df["n4_centro_custo"].astype(str).str.strip()
        + "|"
        + df["valor_conta"].map(norm_valor_chave)
        + "|"
        + df["filial"].astype(str).str.strip().fillna("")
        + "|"
        + df["titulo"].astype(str).str.strip().fillna("")
    )


def montar_mapa_marca() -> dict[str, str]:
    cat = carregar_catalogo_maquinas()
    mp: dict[str, str] = {}
    for _, row in cat.iterrows():
        marca = str(row.get("marca", "")).strip().upper()
        if not marca:
            continue
        for tok in (row["placa"], row["placa_norm"]):
            mp[str(tok).strip().upper()] = marca
            mid = extrair_id_maquina(tok)
            if mid:
                mp[mid] = marca
    return mp


def montar_placas_catalogo() -> set[str]:
    cat = carregar_catalogo_maquinas()
    placas: set[str] = set()
    for p in cat["placa"]:
        placas.add(str(p).strip().upper())
        placas.add(normalizar_placa(str(p)))
    return placas


def is_linha_maquina(row: pd.Series, placas: set[str]) -> bool:
    n4 = str(row.get("n4_centro_custo", "")).strip().upper()
    if extrair_id_maquina(n4):
        return True
    if n4 in placas:
        return True
    if normalizar_placa(n4) in placas:
        return True
    n3 = str(row.get("n3_centro_custo", "")).strip().upper()
    if extrair_id_maquina(n3):
        return True
    n3_txt = str(row.get("n3_centro_custo", "")).upper()
    if "MAQUINAS E EQUIPAMENTOS" in n3_txt or "MÁQUINAS E EQUIPAMENTOS" in n3_txt:
        return bool(n4)
    return False


def resolver_marca(row: pd.Series, mapa_marca: dict[str, str]) -> str:
    existente = str(row.get("MARCA", "")).strip().upper()
    if existente and existente not in {"NAN", ""}:
        return existente
    for col in ("n4_centro_custo", "n3_centro_custo", "titulo"):
        tok = str(row.get(col, "")).strip().upper()
        mid = extrair_id_maquina(tok)
        if mid and mid in mapa_marca:
            return mapa_marca[mid]
        if tok in mapa_marca:
            return mapa_marca[tok]
        norm = normalizar_placa(tok)
        if norm in mapa_marca:
            return mapa_marca[norm]
    return inferir_marca(str(row.get("n4_centro_custo", "")), "", "")


def normalizar_colunas(df: pd.DataFrame, modelo_cols: list[str]) -> pd.DataFrame:
    out = df.copy()
    if "Divisao" not in out.columns and "sdssds" in out.columns:
        out["Divisao"] = out["sdssds"]
    if "Divisao" not in out.columns and "n1_centro_custo" in out.columns:
        out["Divisao"] = out["n1_centro_custo"]
    for col in modelo_cols:
        if col not in out.columns:
            out[col] = ""
    extras = [c for c in out.columns if c not in modelo_cols and c not in ("Arquivo_Fonte", "Tipo_Fonte", "_chave")]
    out = out.drop(columns=extras, errors="ignore")
    return out


def filtrar_periodo(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out["data_nf"] = pd.to_datetime(out["data_nf"], errors="coerce", dayfirst=True)
    mask = (out["data_nf"] >= pd.Timestamp(PERIODO_INI)) & (out["data_nf"] <= pd.Timestamp(PERIODO_FIM))
    return out[mask].copy()


def carregar_fechamento(nome: str, path: Path, placas: set[str]) -> pd.DataFrame:
    if not path.exists():
        print(f"  [aviso] ausente: {path}")
        return pd.DataFrame()
    df = pd.read_excel(path)
    df = filtrar_periodo(df)
    df = df[df.apply(lambda r: is_linha_maquina(r, placas), axis=1)].copy()
    df["Arquivo_Fonte"] = nome
    df["Tipo_Fonte"] = TIPO_FECHAMENTO
    return df


def carregar_diesel_q1() -> pd.DataFrame:
    nome, path = DIESEL_FONTE
    if not path.exists():
        print(f"  [aviso] diesel ausente: {path}")
        return pd.DataFrame()
    df = pd.read_excel(path)
    df = filtrar_periodo(df)
    df["Arquivo_Fonte"] = nome
    df["Tipo_Fonte"] = TIPO_DIESEL
    return df


def normalizar_valor_series(s: pd.Series) -> pd.Series:
    def _one(val):
        if pd.isna(val):
            return val
        if isinstance(val, (int, float)):
            return float(val)
        parsed = parse_numero_br(val)
        return parsed if parsed is not None else val

    return s.map(_one)


def consolidar() -> tuple[pd.DataFrame, pd.DataFrame]:
    if not XLSX_BASE.exists():
        raise FileNotFoundError(f"Base não encontrada: {XLSX_BASE}")

    base_atual = pd.read_excel(XLSX_BASE, sheet_name=SHEET_BASE)
    modelo_cols = [c for c in base_atual.columns if c not in ("Arquivo_Fonte", "Tipo_Fonte")]
    mapa_marca = montar_mapa_marca()
    placas = montar_placas_catalogo()

    print("Carregando fechamentos...")
    partes_fech: list[pd.DataFrame] = []
    for nome, path in FECHAMENTO_SOURCES:
        df = carregar_fechamento(nome, path, placas)
        if not df.empty:
            print(f"  {nome}: {len(df):,} linhas máquina".replace(",", "."))
            partes_fech.append(df)

    fech = pd.concat(partes_fech, ignore_index=True) if partes_fech else pd.DataFrame()
    diesel = carregar_diesel_q1()
    if not diesel.empty:
        print(f"  {DIESEL_FONTE[0]}: {len(diesel):,} linhas".replace(",", "."))

    base = normalizar_colunas(base_atual, modelo_cols)
    prev_tags: dict[str, tuple[str, str]] = {}
    if "Arquivo_Fonte" in base_atual.columns:
        tmp_prev = base_atual.copy()
        tmp_prev["_chave"] = chave_lancamento(tmp_prev)
        for _, pr in tmp_prev.iterrows():
            arq = str(pr.get("Arquivo_Fonte", "")).strip()
            if arq and arq.lower() != "nan":
                prev_tags[pr["_chave"]] = (
                    arq,
                    str(pr.get("Tipo_Fonte", TIPO_SAGI_ATUA)).strip() or TIPO_SAGI_ATUA,
                )

    base["Arquivo_Fonte"] = ""
    base["Tipo_Fonte"] = TIPO_SAGI_ATUA
    base["_chave"] = chave_lancamento(base)

    mapa_fonte: dict[str, str] = {}
    if not fech.empty:
        fech["_chave"] = chave_lancamento(fech)
        for _, row in fech.iterrows():
            mapa_fonte[row["_chave"]] = row["Arquivo_Fonte"]

    matched = 0
    for i, row in base.iterrows():
        k = row["_chave"]
        if k in mapa_fonte:
            base.at[i, "Arquivo_Fonte"] = mapa_fonte[k]
            base.at[i, "Tipo_Fonte"] = TIPO_FECHAMENTO
            matched += 1
        elif k in prev_tags:
            base.at[i, "Arquivo_Fonte"], base.at[i, "Tipo_Fonte"] = prev_tags[k]
        else:
            base.at[i, "Arquivo_Fonte"] = "relatorio_marcas_maquinas_base.xlsx (legado SAGI/Atua)"

    diesel_nome = DIESEL_FONTE[0]
    mask_diesel = (
        base["observacao"].astype(str).str.contains("Custo combust", case=False, na=False)
        | base["titulo"].astype(str).str.match(r"COMB_\d{2}_2026", na=False)
    )
    base.loc[mask_diesel, "Arquivo_Fonte"] = diesel_nome
    base.loc[mask_diesel, "Tipo_Fonte"] = TIPO_DIESEL

    print(f"Linhas existentes vinculadas a fechamento: {matched:,}".replace(",", "."))

    chaves_base = set(base["_chave"])
    novas: list[pd.DataFrame] = []

    if not fech.empty:
        faltantes = fech[~fech["_chave"].isin(chaves_base)].copy()
        if not faltantes.empty:
            print(f"Novas do fechamento: {len(faltantes):,}".replace(",", "."))
            novas.append(faltantes)

    if not diesel.empty:
        diesel["_chave"] = chave_lancamento(diesel)
        diesel_novos = diesel[~diesel["_chave"].isin(chaves_base)].copy()
        if not diesel_novos.empty:
            print(f"Novas diesel Q1/2026: {len(diesel_novos):,}".replace(",", "."))
            novas.append(diesel_novos)
        chaves_base |= set(diesel["_chave"])

    if novas:
        add = pd.concat(novas, ignore_index=True)
        add = normalizar_colunas(add, modelo_cols)
        add["MARCA"] = add.apply(lambda r: resolver_marca(r, mapa_marca), axis=1)
        base = pd.concat([base.drop(columns=["_chave"]), add.drop(columns=["_chave"], errors="ignore")], ignore_index=True)
    else:
        base = base.drop(columns=["_chave"])

    base["data_nf"] = pd.to_datetime(base["data_nf"], errors="coerce", dayfirst=True)
    for col in ("valor_conta", "valor_nf", "valor_pago", "Valor Oficial"):
        if col in base.columns:
            base[col] = normalizar_valor_series(base[col])
    base = base.sort_values(["data_nf", "n4_centro_custo", "cod_conta"]).reset_index(drop=True)

    manifesto_rows = []
    for arquivo, sub in base.groupby("Arquivo_Fonte", dropna=False):
        vals = sub["valor_conta"].map(lambda v: abs(parse_numero_br(v) or pd.to_numeric(v, errors="coerce") or 0))
        manifesto_rows.append({
            "Arquivo_Fonte": arquivo,
            "Tipo_Fonte": sub["Tipo_Fonte"].iloc[0] if "Tipo_Fonte" in sub.columns else "",
            "Linhas": len(sub),
            "Soma_abs_valor_conta": float(vals.sum()),
            "Periodo_min": sub["data_nf"].min(),
            "Periodo_max": sub["data_nf"].max(),
        })
    manifesto = pd.DataFrame(manifesto_rows).sort_values("Linhas", ascending=False)

    return base, manifesto


def salvar(base: pd.DataFrame, manifesto: pd.DataFrame) -> Path:
    from openpyxl import load_workbook

    out_cols = [c for c in base.columns if c != "_chave"]
    try:
        wb = load_workbook(XLSX_BASE)
    except FileNotFoundError:
        wb = None

    if wb is None:
        with pd.ExcelWriter(XLSX_BASE, engine="openpyxl") as writer:
            base[out_cols].to_excel(writer, sheet_name=SHEET_BASE, index=False)
            manifesto.to_excel(writer, sheet_name=SHEET_MANIFESTO, index=False)
        return XLSX_BASE

    for sheet in (SHEET_BASE, SHEET_MANIFESTO):
        if sheet in wb.sheetnames:
            del wb[sheet]
    wb.save(XLSX_BASE)

    with pd.ExcelWriter(XLSX_BASE, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
        base[out_cols].to_excel(writer, sheet_name=SHEET_BASE, index=False)
        manifesto.to_excel(writer, sheet_name=SHEET_MANIFESTO, index=False)
    return XLSX_BASE


def main() -> None:
    print(f"Consolidando -> {XLSX_BASE}")
    base, manifesto = consolidar()
    path = salvar(base, manifesto)
    print(f"\nBase atualizada: {len(base):,} linhas".replace(",", "."))
    print(f"Manifesto ({len(manifesto)} fontes):")
    for _, r in manifesto.iterrows():
        print(f"  - {r['Arquivo_Fonte']}: {int(r['Linhas']):,} linhas | R$ {r['Soma_abs_valor_conta']:,.2f}".replace(",", "."))
    print(f"\nArquivo: {path}")


if __name__ == "__main__":
    main()
