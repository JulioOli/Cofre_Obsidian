"""Validate dynamic April closing workbook: tables + formulas."""
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SAIDA = ROOT / "02-Referencias" / "04_2026_Abr_Rel_FOPA.xlsx"

wb = openpyxl.load_workbook(SAIDA)
print(f"Abas: {wb.sheetnames}\n")

ws = wb["GERAL"]
print("== GERAL ==")
print(f"  Dim: {ws.calculate_dimension()}")
print(f"  Tabelas: {list(ws.tables.keys())}")
for name in ws.tables:
    tbl = ws.tables[name]
    print(f"    {name}: ref={tbl!r}")
print("\n  Linha 1 (título + totais):")
for col, cel in enumerate(ws[1], start=1):
    if cel.value is not None:
        print(f"    {chr(64+col)}1: {cel.value!r}")
print("\n  Linha 2 (cabeçalhos):")
for col, cel in enumerate(ws[2], start=1):
    if cel.value is not None:
        print(f"    {chr(64+col)}2: {cel.value!r}")
print("\n  Linhas 3-5 (amostra):")
for r in (3, 4, 5):
    vals = []
    for cel in ws[r]:
        if cel.value is not None:
            vals.append(f"{cel.coordinate}={cel.value!r}")
    print(f"    {' | '.join(vals[:5])}{' ...' if len(vals) > 5 else ''}")

for nome in ("PILARES", "BRACOFER", "SELETIVA PRUDENTE"):
    ws = wb[nome]
    print(f"\n== {nome} ==")
    print(f"  Dim: {ws.calculate_dimension()}")
    print("  Fórmulas (apenas células com '='):")
    for row in ws.iter_rows(min_row=1, max_row=3):
        for cel in row:
            if isinstance(cel.value, str) and cel.value.startswith("="):
                print(f"    {cel.coordinate}: {cel.value}")
            elif cel.value not in (None, ""):
                print(f"    {cel.coordinate}: (valor) {cel.value!r}")
