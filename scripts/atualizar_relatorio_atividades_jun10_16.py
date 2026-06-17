"""Atualiza Relatorio_de_Atividades.xlsx para 10-16/06/2026."""
from __future__ import annotations

import shutil
from copy import copy
from datetime import datetime, time
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "02-Referencias" / "Meus_Dados" / "Relatorio_de_Atividades.xlsx"


def copy_row_style(ws, src_row: int, dst_row: int, max_col: int = 7) -> None:
    for col in range(1, max_col + 1):
        src = ws.cell(src_row, col)
        dst = ws.cell(dst_row, col)
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
        dst.number_format = src.number_format


def set_activity(ws, row: int, day: int, empresa: str, resumo: str, desc: str, ini: time, fim: time) -> None:
    ws.cell(row, 1, f"=DATE(2026,6,{day})")
    ws.cell(row, 2, empresa)
    ws.cell(row, 3, resumo)
    ws.cell(row, 4, desc)
    ws.cell(row, 5, ini)
    ws.cell(row, 6, fim)
    ws.cell(row, 7, f"=F{row}-E{row}")


def add_banco_row(ws_bh, row: int, day: int, ini: time, fim: time, prev_row: int) -> None:
    ws_bh.cell(row, 1, f"=DATE(2026,6,{day})")
    ws_bh.cell(row, 2, ini)
    ws_bh.cell(row, 3, fim)
    ws_bh.cell(row, 4, f"=IF(C{row}-B{row}>TIME(6,0,0), (C{row}-B{row})-TIME(1,0,0), C{row}-B{row})")
    ws_bh.cell(row, 5, f'=IF(OR(D{row}<=0, B{row}="", C{row}=""), 0, TIME(6,0,0))')
    ws_bh.cell(row, 6, f"=D{row}-E{row}")
    ws_bh.cell(row, 7, time(0, 0))
    ws_bh.cell(row, 8, f"=F{row}-G{row}")
    ws_bh.cell(row, 9, f"=I{prev_row}+H{row}")


def main() -> None:
    wb = load_workbook(XLSX)
    ws = wb["Atividades"]
    ws_bh = wb["Banco de Horas"]

    style_row = 559
    extra_rows = [560, 561, 562, 563, 564, 565, 566, 567, 568]
    for r in extra_rows:
        copy_row_style(ws, style_row, r)

    for r in range(635, 642):
        copy_row_style(ws, style_row, r)

    # 10/06 — R559 já existe (07:52-09:36)
    set_activity(
        ws,
        635,
        10,
        "G3S 10",
        "Fechamento",
        "Conferência e ajuste de base para fechamento de custos de combustível e folha "
        "(conferir_folha05 e relatórios custoComb_mes).",
        time(9, 36),
        time(12, 0),
    )
    set_activity(
        ws,
        560,
        10,
        "G3S 10",
        "Fechamento",
        "Preparação dos arquivos de fechamento de combustível por placa "
        "(relatorio_custoComb_mes_fechamento).",
        time(13, 0),
        time(15, 8),
    )

    # 11/06
    set_activity(
        ws,
        636,
        11,
        "G3S 10",
        "Suporte áreas",
        "Análise comparativa Liebherr × Hyundai — refinamento de metodologia e consolidação "
        "de insumos para relatório executivo (notebook analise_comparativa_liebherr_hyundai).",
        time(7, 55),
        time(10, 0),
    )
    set_activity(
        ws,
        561,
        11,
        "G3S 10",
        "Suporte áreas",
        "Levantamento entre empresas e quantidade de máquinas por filial "
        "(notas Zettelkasten e outputs de apoio).",
        time(10, 0),
        time(12, 0),
    )
    set_activity(
        ws,
        562,
        11,
        "G3S 10",
        "Fechamento",
        "Tratamento dos relatórios custoComb placas jan-mar (prep e fechamento).",
        time(13, 0),
        time(15, 6),
    )

    # 12/06
    set_activity(
        ws,
        637,
        12,
        "G3S 10",
        "Suporte áreas",
        "Tratamento de dados de consumo de diesel e geração de análises de contratos para máquinas "
        "(scripts medias_consumo_diesel e gerar_analise_contratos_diesel_maquinas).",
        time(7, 52),
        time(10, 30),
    )
    set_activity(
        ws,
        563,
        12,
        "G3S 10",
        "Suporte áreas",
        "Auditoria BCBI e consolidação do relatório de marcas de máquinas "
        "(notebook auditoria_bcbi e relatorio_marcas_maquinas).",
        time(10, 30),
        time(12, 0),
    )
    set_activity(
        ws,
        564,
        12,
        "G3S 10",
        "Fechamento",
        "Normalização Supply e ajustes nos notebooks/scripts de fechamento "
        "(normalizar_supply.ipynb e fechamento_excel.py).",
        time(13, 0),
        time(15, 8),
    )

    # 13/06 (sábado — carga reduzida)
    set_activity(
        ws,
        638,
        13,
        "G3S 10",
        "Organização técnica",
        "Organização e revisão técnica de artefatos de análise e documentação de apoio do projeto.",
        time(9, 0),
        time(12, 0),
    )

    # 14/06 (domingo — carga reduzida)
    set_activity(
        ws,
        639,
        14,
        "G3S 10",
        "Organização técnica",
        "Consolidação de pendências e preparação de material para continuidade "
        "das entregas da semana seguinte.",
        time(10, 0),
        time(12, 0),
    )

    # 15/06
    set_activity(
        ws,
        640,
        15,
        "G3S 10",
        "BCBI Seletiva",
        "Operação assistida do sistema de BI da seletiva, acompanhamento e checagem de erros no sistema.",
        time(7, 48),
        time(9, 0),
    )
    set_activity(
        ws,
        565,
        15,
        "G3S 10",
        "Suporte áreas",
        "Atualização do dashboard comparativo Liebherr × Hyundai, sincronização de figuras "
        "e ajustes de scripts de exportação (export_dashboard_excel e sincronizar_figuras).",
        time(9, 0),
        time(11, 30),
    )
    set_activity(
        ws,
        566,
        15,
        "G3S 10",
        "Suporte áreas",
        "Consolidação da base de máquinas para fechamento e atualização das análises de contratos diesel "
        "(consolidar_base_maquinas_fechamento e analise_contratos_diesel_maquinas).",
        time(12, 30),
        time(15, 6),
    )

    # 16/06
    set_activity(
        ws,
        641,
        16,
        "G3S 10",
        "Governança de dados",
        "Atualização de referências de Plano de Contas e revisão das notas de CC "
        "(Divisões, Filiais e Departamentos).",
        time(7, 52),
        time(10, 30),
    )
    set_activity(
        ws,
        567,
        16,
        "G3S 10",
        "Governança de dados",
        "Revisão de apontamentos de auditoria analítica BCBI para suporte à classificação contábil.",
        time(10, 30),
        time(12, 0),
    )
    set_activity(
        ws,
        568,
        16,
        "G3S 10",
        "Suporte áreas",
        "Atualização do relatório de atividades e revisão das entregas da semana.",
        time(13, 0),
        time(15, 0),
    )

    # Banco de Horas — copiar estilo da linha 132
    bh_template = 132
    banco_days = [
        (133, 11, time(7, 55), time(15, 6), 132),
        (134, 12, time(7, 52), time(15, 8), 133),
        (135, 13, time(9, 0), time(12, 0), 134),
        (136, 14, time(10, 0), time(12, 0), 135),
        (137, 15, time(7, 48), time(15, 6), 136),
        (138, 16, time(7, 52), time(15, 0), 137),
    ]
    for row, day, ini, fim, prev in banco_days:
        copy_row_style(ws_bh, bh_template, row, max_col=9)
        add_banco_row(ws_bh, row, day, ini, fim, prev)

    wb.save(XLSX)
    print(f"Salvo: {XLSX}")

    # Validação
    wb2 = load_workbook(XLSX, data_only=False)
    ws2 = wb2["Atividades"]
    print("\n=== Validação linhas preenchidas ===")
    for r in [559] + extra_rows + list(range(635, 642)):
        vals = [ws2.cell(r, c).value for c in range(1, 8)]
        if any(vals):
            print(f"R{r}: A={vals[0]!r} E={vals[4]!r} F={vals[5]!r} G={vals[6]!r}")

    ws_bh2 = wb2["Banco de Horas"]
    print("\n=== Banco de Horas novas linhas ===")
    for r in range(133, 139):
        vals = [ws_bh2.cell(r, c).value for c in range(1, 10)]
        print(f"R{r}: {vals}")


if __name__ == "__main__":
    main()
