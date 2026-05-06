"""
Preenche CONTROLE DE PONTO.xlsx a partir de Atividades.csv (INICIAL/FINAL por bloco),
agrupando o dia em dois períodos (antes e depois do almoço).
"""
from __future__ import annotations

import csv
import re
import sys
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CSV = ROOT / "02-Referencias" / "Atividades.csv"
DEFAULT_XLSX = ROOT / "02-Referencias" / "CONTROLE DE PONTO.xlsx"


def _detect_encoding(path: Path) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            path.read_text(encoding=enc)[:50_000]
            return enc
        except UnicodeDecodeError:
            continue
    return "latin-1"


def parse_time(s: str) -> time | None:
    s = (s or "").strip()
    if not s:
        return None
    m = re.match(r"^(\d{1,2}):(\d{2})(?::(\d{2}))?$", s)
    if not m:
        return None
    h, mi = int(m.group(1)), int(m.group(2))
    return time(h, mi)


def to_minutes(t: time) -> int:
    return t.hour * 60 + t.minute


def from_minutes(m: int) -> time:
    m = max(0, m)
    return time(m // 60, m % 60)


@dataclass(frozen=True)
class Block:
    start: time
    end: time


def merge_touching(blocks: list[Block]) -> list[Block]:
    if not blocks:
        return []
    b = sorted(blocks, key=lambda x: to_minutes(x.start))
    out: list[Block] = [b[0]]
    for cur in b[1:]:
        last = out[-1]
        if to_minutes(cur.start) <= to_minutes(last.end):
            if to_minutes(cur.end) > to_minutes(last.end):
                out[-1] = Block(last.start, cur.end)
        else:
            out.append(cur)
    return out


def split_morning_afternoon(blocks: list[Block]) -> tuple[Block | None, Block | None]:
    """
    Retorna (manhã, tarde) como blocos agregados (entrada/saída do período).
    """
    if not blocks:
        return None, None
    blocks = merge_touching(blocks)
    if len(blocks) == 1:
        b = blocks[0]
        sm, em = to_minutes(b.start), to_minutes(b.end)
        # Jornada contínua cobrindo almoço: divide em 12:00 / 13:00
        if sm < 12 * 60 and em > 13 * 60 and (em - sm) >= 5 * 60:
            return Block(b.start, time(12, 0)), Block(time(13, 0), b.end)
        return b, None

    best_i: int | None = None
    best_score = -1.0
    for i in range(len(blocks) - 1):
        gap = to_minutes(blocks[i + 1].start) - to_minutes(blocks[i].end)
        if gap < 15:
            continue
        end_m = to_minutes(blocks[i].end)
        start_next = to_minutes(blocks[i + 1].start)
        # Janela típica de almoço
        if end_m < 10 * 60 + 30 or start_next > 15 * 60:
            continue
        score = gap
        if 11 * 60 + 0 <= end_m <= 14 * 60 + 30:
            score += 60
        if 12 * 60 <= start_next <= 15 * 60:
            score += 60
        if score > best_score:
            best_score = score
            best_i = i

    if best_i is None:
        return blocks[0], None

    morning = Block(blocks[0].start, blocks[best_i].end)
    afternoon = Block(blocks[best_i + 1].start, blocks[-1].end)
    return morning, afternoon


def load_april_2026_blocks(path: Path) -> dict[int, list[Block]]:
    enc = _detect_encoding(path)
    by_day: dict[int, list[Block]] = {}
    with open(path, newline="", encoding=enc) as f:
        r = csv.DictReader(f)
        for row in r:
            ds = (row.get("DATA") or "").strip()
            if not ds:
                continue
            try:
                dt = datetime.strptime(ds, "%m/%d/%Y").date()
            except ValueError:
                continue
            if dt.year != 2026 or dt.month != 4:
                continue
            a = parse_time(row.get("INICIAL", "") or "")
            b = parse_time(row.get("FINAL", "") or "")
            if a is None or b is None:
                continue
            if to_minutes(b) < to_minutes(a):
                continue
            by_day.setdefault(dt.day, []).append(Block(a, b))
    return by_day


def fill_xlsx(
    xlsx_path: Path,
    by_day: dict[int, list[Block]],
    *,
    nome: str,
    setor: str,
    mes_label: str,
    sheet_title: str,
) -> None:
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active
    ws.title = sheet_title
    ws["A2"].value = f"NOME: {nome}"
    ws["A3"].value = f"SETOR: {setor}"
    ws["A5"].value = f"MÊS: {mes_label}"

    for day in range(1, 32):
        rr = 7 + day
        for col in range(2, 6):
            ws.cell(rr, col).value = None

    for day, raw_blocks in sorted(by_day.items()):
        morn, aft = split_morning_afternoon(raw_blocks)
        rr = 7 + day
        if morn:
            ws.cell(rr, 2).value = morn.start
            ws.cell(rr, 3).value = morn.end
        if aft:
            ws.cell(rr, 4).value = aft.start
            ws.cell(rr, 5).value = aft.end

    wb.save(xlsx_path)


def main() -> int:
    nome = "Julio Oliveira Santana"
    setor = "Qualidade"
    mes_label = "Abril"
    csv_path = DEFAULT_CSV
    xlsx_path = DEFAULT_XLSX
    if len(sys.argv) > 1:
        csv_path = Path(sys.argv[1])
    if len(sys.argv) > 2:
        xlsx_path = Path(sys.argv[2])

    by_day = load_april_2026_blocks(csv_path)
    fill_xlsx(
        xlsx_path,
        by_day,
        nome=nome,
        setor=setor,
        mes_label=mes_label,
        sheet_title="Ponto Abr 2026",
    )
    print(f"Atualizado: {xlsx_path} ({len(by_day)} dias com atividades em abril/2026)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
