"""
Gera planilha de análise de máquinas por contrato e consumo de diesel (interno x externo).

Fontes:
  - 02-Referencias/Custos-Maquinas/relatorio_marcas_maquinas.xlsx
      (abas CC_Maquinas_2026, entre_empresas_abril)
  - 02-Referencias/Custos-Maquinas/relatorio_marcas_maquinas_base.xlsx
      (aba base_maquinas_2025_2026 — lançamentos de fechamento)

Saída:
  - 02-Referencias/Custos-Maquinas/analise_contratos_diesel_maquinas.xlsx
    • Maquinas_por_Contrato
    • Diesel_Interno_Externo

Recorte padrão: jan/2025 a mar/2026 | somente escavadeiras Hyundai / Liebherr (EHH/EHL).

Médias de consumo operacional (planilha de abastecimento interno): ver script separado
  scripts/tratar_medias_consumo_diesel.py → medias_consumo_diesel_tratado.xlsx
"""

from __future__ import annotations

import re
import unicodedata
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
CUSTOS = ROOT / "02-Referencias" / "Custos-Maquinas"
XLSX_IN = CUSTOS / "relatorio_marcas_maquinas.xlsx"
XLSX_BASE = CUSTOS / "relatorio_marcas_maquinas_base.xlsx"
XLSX_OUT = CUSTOS / "analise_contratos_diesel_maquinas.xlsx"
XLSX_MEDIA = CUSTOS / "Media_de_Consumo_de_Diesel_das_Maquinas.xlsx"

SHEET_CC = "CC_Maquinas_2026"
SHEET_MEDIA = "Media Maquinas 2026"
SHEET_EE = "entre_empresas_abril"
SHEET_BASE = "base_maquinas_2025_2026"

CONTAS_DIESEL = {
    "7.1.4": ("Externo", "Diesel — Posto (externo)"),
    "7.1.22": ("Interno", "Diesel — Interno"),
    "7.1.23": ("Interno", "Diesel — Estoque (abastecimento interno)"),
}

# Recorte alinhado às análises Liebherr × Hyundai (analise_comparativa_liebherr_hyundai.ipynb)
PERIODO_INI = date(2025, 1, 1)
PERIODO_FIM = date(2026, 3, 31)
MARCAS_ALVO = frozenset({"HYUNDAI", "LIEBHERR"})

PARQUE_ESCAVADEIRAS: frozenset[str] = frozenset(
    {
        # Liebherr (7)
        "EHL0013", "EHL0014", "EHL0028", "EHL0041", "EHL0042", "EHL0043", "EHL0070",
        # Hyundai (16)
        "EHH0002", "EHH0003", "EHH0004", "EHH0006", "EHH0007", "EHH0008", "EHH0009",
        "EHH0015", "EHH0019", "EHH0036", "EHH0037", "EHH0038", "EHH0039", "EHH0040",
        "EHH0044", "EHH0046",
    }
)

# Coluna G (MÉDIA CONSUMO L/H): valores muito discrepantes são zerados para não distorcer a média geral
LH_MIN_COL_G = 3.0
LH_MAX_COL_G = 40.0

FONT_HEADER = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
FILL_HEADER = PatternFill("solid", fgColor="305496")
FMT_BRL = 'R$ #,##0.00;[Red]-R$ #,##0.00;""'


def _norm_txt(val) -> str:
    if pd.isna(val):
        return ""
    return str(val).strip()


def limpar_localizacao(val) -> str:
    """Trata '-' e '?' do CC como vazio para permitir fallback da aba entre_empresas."""
    s = _norm_txt(val)
    return "" if s in {"-", "?", "NAN"} else s


def normalizar_placa(placa: str) -> str:
    parts = re.split(r"[/\s]+", _norm_txt(placa))
    token = next((p for p in parts if p and re.match(r"^[A-Z0-9]", p, re.I)), "")
    return token.upper()


def is_placa_valida(placa: str) -> bool:
    s = _norm_txt(placa).upper()
    if not s or s in {"-", "?", "NAN"}:
        return False
    if "UNIDADES" in s or s == "GRIMAL":
        return False
    return bool(re.match(r"^[A-Z0-9]{3,}", s))


def is_tipo_maquina(tipo: str) -> bool:
    norm = (
        unicodedata.normalize("NFKD", str(tipo or ""))
        .encode("ascii", "ignore")
        .decode("ascii")
        .upper()
    )
    return any(k in norm for k in ("MAQUINA", "PRENSA", "BALANCA"))


def to_float(val) -> float | None:
    if pd.isna(val):
        return None
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def parse_numero_br(val) -> float | None:
    """Converte número em formato BR (1.234,56) ou misto do Excel."""
    if pd.isna(val):
        return None
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().upper().replace("R$", "").replace(" ", "")
    if not s or s in {"-", "NAN"}:
        return None
    # Corrige typos como 2.3484,28
    m_typo = re.match(r"^(\d+)\.(\d+),(\d+)$", s)
    if m_typo:
        s = f"{m_typo.group(1)}{m_typo.group(2)}.{m_typo.group(3)}"
    elif "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def parse_horas_trabalhadas(val) -> float | None:
    """Extrai horas decimais de textos como '(31) = 31h' ou '0,0'."""
    if pd.isna(val):
        return None
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip()
    if not s:
        return None
    m = re.search(r"\((\d+(?:[.,]\d+)?)\)", s)
    if m:
        return parse_numero_br(m.group(1))
    m = re.search(r"(\d+(?:[.,]\d+)?)\s*h", s, re.I)
    if m:
        return parse_numero_br(m.group(1))
    return parse_numero_br(s)


def horimetro_suspeito(val) -> bool:
    n = parse_numero_br(val)
    return n is not None and n >= 9000


def media_col_g_discrepante(media_g: float | None) -> tuple[bool, str]:
    """Coluna G muito discrepante (outlier) — será zerada no cálculo da média geral."""
    if media_g is None:
        return True, "coluna G vazia"
    if media_g < LH_MIN_COL_G:
        return True, f"G < {LH_MIN_COL_G} L/h"
    if media_g > LH_MAX_COL_G:
        return True, f"G > {LH_MAX_COL_G} L/h"
    return False, ""


def tratar_media_col_g(media_g: float | None) -> tuple[float | None, bool, str]:
    """Zera coluna G quando discrepante; mantém o abastecimento no detalhe."""
    discrepante, motivo = media_col_g_discrepante(media_g)
    if discrepante:
        return 0.0, True, motivo
    return media_g, False, ""


def is_escavadeira_hyundai_liebherr(placa_norm: str, marca: str = "", implemento: str = "") -> bool:
    """Escavadeira Hyundai/Liebherr — parque oficial (7 Liebherr + 16 Hyundai)."""
    return _norm_txt(placa_norm).upper() in PARQUE_ESCAVADEIRAS


def filtrar_periodo(df: pd.DataFrame, col_data: str = "data_nf") -> pd.DataFrame:
    out = df.copy()
    out[col_data] = pd.to_datetime(out[col_data], errors="coerce")
    mask = (out[col_data] >= pd.Timestamp(PERIODO_INI)) & (out[col_data] <= pd.Timestamp(PERIODO_FIM))
    return out[mask].copy()


def inferir_marca(placa: str, marca_cc: str, implemento: str) -> str:
    marca = _norm_txt(marca_cc)
    if marca:
        return marca
    plc = _norm_txt(placa).upper()
    impl = _norm_txt(implemento).upper()
    if plc.startswith("EHL") or "LIEBHERR" in impl:
        return "LIEBHERR"
    if plc.startswith("EHH"):
        return "HYUNDAI"
    return ""


def classificar_contrato(
    localizacao: str,
    primeira_perna: str = "",
    segunda_perna: str = "",
    n3_centro: str = "",
) -> tuple[str, str, str]:
    """Retorna (grupo_cliente, contrato_detalhe, tipo_operacao)."""
    loc = _norm_txt(localizacao).upper()
    p1 = _norm_txt(primeira_perna).upper()
    p2 = _norm_txt(segunda_perna).upper()
    n3 = _norm_txt(n3_centro).upper()

    # Prioridade: CC / localização explícita de contrato externo
    if "TUPY" in n3 or "JOINVILLE" in loc:
        return "Tupy", "Tupy — Joinville/SC", "Contrato externo"
    if "GERDAU" in n3 or "GERDAU" in loc:
        return "Gerdau", "Gerdau", "Contrato externo"
    if "ARCELOR" in n3 or "ARCELOR" in loc or "PIRACICABA" in loc:
        if "RESENDE" in loc or "RESENDE" in n3:
            return "Arcelor", "Arcelor — Resende/RJ", "Contrato externo"
        if "BARRA MANSA" in loc or "BARRA MANSA" in n3:
            return "Arcelor", "Arcelor — Barra Mansa/RJ", "Contrato externo"
        if "PIRACICABA" in loc or "PIRACICABA" in n3:
            return "Arcelor", "Arcelor — Piracicaba/SP", "Contrato externo"
        if "IRACEMA" in loc or "IRACEMA" in n3:
            return "Arcelor", "Arcelor — Iracemápolis/SP", "Contrato externo"
        return "Arcelor", "Arcelor — (planta não especificada)", "Contrato externo"
    if "RESENDE" in loc:
        return "Arcelor", "Arcelor — Resende/RJ", "Contrato externo"
    if "BARRA MANSA" in loc:
        return "Arcelor", "Arcelor — Barra Mansa/RJ", "Contrato externo"

    if "G&S PARA CLIENTE" in p2 or "RSE DIRETO PARA CLIENTE" in p1:
        return "Contrato externo (outro)", loc or "Cliente externo", "Contrato externo"

    if loc == "RSE" or loc.startswith("RSE "):
        return "RSE", "RSE — pátio / manutenção", "Uso interno grupo"

    if "G3S" in p2 or "G3S" in p1 or "SELETIVA" in loc:
        cidade = loc
        for prefix in (
            "SP - PRUDENTE - SUCATA",
            "SP - PRUDENTE - RESERVA",
            "SP - SELETIVA",
            "MS - CAMPO GRANDE",
            "MS - DOURADOS",
            "PR - LONDRINA",
            "PR - MARINGÁ (CIDADE ALTA)",
            "PR - MARINGÁ",
            "G3S - FILIAL",
        ):
            if loc.startswith(prefix.upper()) or prefix.upper() in loc:
                cidade = prefix
                break
        if cidade in {"-", "?", ""}:
            return "Seletiva (interno G3S)", "Seletiva — (localização pendente)", "Intercompany G3S"
        return "Seletiva (interno G3S)", f"Seletiva — {cidade}", "Intercompany G3S"

    if "G&S" in p1:
        return "G&S (operacional)", loc or "G&S", "Operação G&S"

    # Máquinas só no CC_Maquinas (sem pernas) — inferir pela localização física
    seletiva_map = [
        ("PRUDENTE", "Seletiva — SP Prudente"),
        ("CAMPO GRANDE", "Seletiva — MS Campo Grande"),
        ("DOURADOS", "Seletiva — MS Dourados"),
        ("LONDRINA", "Seletiva — PR Londrina"),
        ("CIDADE ALTA", "Seletiva — PR Maringá Cidade Alta"),
        ("MARING", "Seletiva — PR Maringá"),
        ("SELETIVA", "Seletiva — SP"),
        ("G3S - FILIAL", "Seletiva — G3S Filial 5"),
    ]
    for chave, nome in seletiva_map:
        if chave in loc:
            return "Seletiva (interno G3S)", nome, "Intercompany G3S"

    if loc in {"-", "?", ""} or loc.isdigit():
        return "RSE / estoque", "RSE — sem localização definida", "Uso interno grupo"

    if "LOJA" in loc or "OFICINA" in loc or "MANUTEN" in loc:
        return "RSE", "RSE — oficina / loja", "Uso interno grupo"

    return "A classificar", loc or "—", "Revisar"


def classificar_contrato_por_cc(n1: str, n3: str) -> tuple[str, str]:
    return classificar_contrato(localizacao="", n3_centro=n3, primeira_perna="", segunda_perna="")


def carregar_catalogo_maquinas() -> pd.DataFrame:
    if not XLSX_IN.exists():
        raise FileNotFoundError(f"Arquivo não encontrado: {XLSX_IN}")

    cc = pd.read_excel(XLSX_IN, sheet_name=SHEET_CC)
    cc = cc[cc["PLACA"].apply(is_placa_valida) & cc["TIPO"].apply(is_tipo_maquina)].copy()

    ee = pd.read_excel(XLSX_IN, sheet_name=SHEET_EE)
    ee = ee[ee["PLACA"].apply(is_placa_valida) & ee["TIPO"].apply(is_tipo_maquina)].copy()

    cc["placa_norm"] = cc["PLACA"].apply(normalizar_placa)
    ee["placa_norm"] = ee["PLACA"].apply(normalizar_placa)

    cc_idx = {r["placa_norm"]: r for _, r in cc.iterrows()}
    ee_idx = {r["placa_norm"]: r for _, r in ee.iterrows()}
    chaves = sorted(set(cc_idx) | set(ee_idx))

    col_contrato_ee = "N.º CONTRATO"
    col_deprec_ee = "VALOR DA LOCAÇÃO (DEPRECIAÇÃO)"

    linhas = []
    for chave in chaves:
        c = cc_idx.get(chave)
        e = ee_idx.get(chave)

        placa = c["PLACA"] if c is not None else e["PLACA"]
        localizacao = limpar_localizacao(c["LOCALIZAÇÃO"] if c is not None else "")
        if not localizacao and e is not None:
            localizacao = limpar_localizacao(e["LOCALIZAÇÃO"])

        implemento = _norm_txt(c["IMPLEMENTO"] if c is not None else "")
        if not implemento and e is not None:
            implemento = _norm_txt(e["IMPLEMENTO"])

        primeira = _norm_txt(e["PRIMEIRA \"PERNA\""]) if e is not None else ""
        segunda = _norm_txt(e["SEGUNDA \"PERNA\""]) if e is not None else ""
        terceira = _norm_txt(e["TERCEIRA \"PERNA\""]) if e is not None else ""

        grupo, contrato, tipo_op = classificar_contrato(localizacao, primeira, segunda)

        deprec = None
        if e is not None:
            deprec = to_float(e.get(col_deprec_ee))
        if deprec is None and c is not None:
            deprec = to_float(c.get("DEPRECIAÇÃO COMO ALUGUEL")) or to_float(c.get("DEPRECIAÇÃO"))

        linhas.append(
            {
                "placa": placa,
                "placa_norm": chave,
                "tipo": _norm_txt(c["TIPO"] if c is not None else e["TIPO"]),
                "marca": inferir_marca(placa, c["MARCA"] if c is not None else "", implemento),
                "modelo": _norm_txt(c["MODELO"] if c is not None else ""),
                "implemento": implemento,
                "localizacao": localizacao,
                "grupo_cliente": grupo,
                "contrato_detalhe": contrato,
                "tipo_operacao": tipo_op,
                "n_contrato_rse": _norm_txt(e[col_contrato_ee]) if e is not None else "",
                "primeira_perna": primeira,
                "segunda_perna": segunda,
                "terceira_perna": terceira,
                "locacao_ou_servico": _norm_txt(e["LOCAÇÃO OU SERVIÇO"]) if e is not None else "LOCAÇÃO",
                "depreciacao_mensal": deprec,
                "centro_custo_sagi": _norm_txt(c["CENTRO DE CUSTO"] if c is not None else ""),
                "status": _norm_txt(c["STATUS DO VEÍCULO"] if c is not None else e.get("STATUS DO VEÍCULO", "")),
                "fonte": "+".join(
                    x for x in ("CC_Maquinas_2026" if c is not None else "", "entre_empresas_abril" if e is not None else "") if x
                ),
            }
        )

    cat = pd.DataFrame(linhas)
    cat = cat[
        cat.apply(
            lambda r: is_escavadeira_hyundai_liebherr(r["placa_norm"], r["marca"], r["implemento"]),
            axis=1,
        )
    ].copy()
    return cat.sort_values(["grupo_cliente", "contrato_detalhe", "placa"])


def montar_resumo_contratos(cat: pd.DataFrame) -> pd.DataFrame:
    g = (
        cat.groupby(["grupo_cliente", "contrato_detalhe", "tipo_operacao"], dropna=False)
        .agg(
            qtd_maquinas=("placa", "count"),
            depreciacao_mensal_total=("depreciacao_mensal", lambda s: pd.to_numeric(s, errors="coerce").sum()),
        )
        .reset_index()
        .sort_values(["grupo_cliente", "contrato_detalhe"])
    )
    return g


def extrair_placa_n4(n4: str) -> str:
    m = re.search(r"([A-Z]{3}\d{4})", str(n4).upper())
    return m.group(1) if m else ""


def carregar_diesel(cat: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    if not XLSX_BASE.exists():
        raise FileNotFoundError(
            f"Base de lançamentos não encontrada: {XLSX_BASE}\n"
            "Execute o notebook relatorio_marcas_maquinas.ipynb para gerá-la."
        )

    df = pd.read_excel(XLSX_BASE, sheet_name=SHEET_BASE)
    df = df[df["cod_conta"].astype(str).isin(CONTAS_DIESEL)].copy()
    df["valor"] = pd.to_numeric(df["valor_conta"], errors="coerce")
    df["valor_abs"] = df["valor"].abs()
    df["placa"] = df["n4_centro_custo"].apply(extrair_placa_n4)

    placas_alvo = set(cat["placa_norm"])
    df = df[df["placa"].isin(placas_alvo)].copy()
    df = filtrar_periodo(df)

    mapa_placa = cat.set_index("placa_norm")[
        ["grupo_cliente", "contrato_detalhe", "tipo_operacao", "localizacao", "marca"]
    ].to_dict("index")

    def enrich(row):
        placa = row["placa"]
        info = mapa_placa.get(placa, {})
        if info:
            return pd.Series(
                {
                    "marca": info.get("marca", ""),
                    "grupo_cliente": info.get("grupo_cliente", ""),
                    "contrato_detalhe": info.get("contrato_detalhe", ""),
                    "tipo_operacao": info.get("tipo_operacao", ""),
                    "localizacao_catalogo": info.get("localizacao", ""),
                }
            )
        g, c, t = classificar_contrato_por_cc(
            str(row.get("n1_cod_centro_custo", "")),
            str(row.get("n3_centro_custo", "")),
        )
        return pd.Series(
            {
                "marca": "",
                "grupo_cliente": g,
                "contrato_detalhe": c,
                "tipo_operacao": t,
                "localizacao_catalogo": "",
            }
        )

    enrich_cols = df.apply(enrich, axis=1)
    df = pd.concat([df, enrich_cols], axis=1)

    conta_map = {k: v for k, v in CONTAS_DIESEL.items()}
    df["diesel_origem"] = df["cod_conta"].map(lambda c: conta_map.get(str(c), ("", ""))[0])
    df["diesel_conta_desc"] = df["cod_conta"].map(lambda c: conta_map.get(str(c), ("", ""))[1])

    df["data_nf"] = pd.to_datetime(df["data_nf"], errors="coerce")
    df["competencia"] = df["data_nf"].dt.to_period("M").astype(str)

    detalhe = df[
        [
            "competencia",
            "data_nf",
            "placa",
            "marca",
            "grupo_cliente",
            "contrato_detalhe",
            "tipo_operacao",
            "diesel_origem",
            "diesel_conta_desc",
            "cod_conta",
            "conta",
            "valor",
            "valor_abs",
            "n1_centro_custo",
            "n3_centro_custo",
            "n4_centro_custo",
            "filial",
            "credor_forn_cli_func",
            "observacao",
        ]
    ]
    detalhe = detalhe.sort_values(["grupo_cliente", "diesel_origem", "competencia", "placa"])

    resumo = (
        df.groupby(["grupo_cliente", "contrato_detalhe", "diesel_origem", "diesel_conta_desc"], dropna=False)
        .agg(
            qtd_lancamentos=("valor", "count"),
            valor_total=("valor", "sum"),
            valor_abs_total=("valor_abs", "sum"),
        )
        .reset_index()
        .sort_values(["grupo_cliente", "diesel_origem", "contrato_detalhe"])
    )

    resumo_placa = (
        df.groupby(["placa", "marca", "grupo_cliente", "contrato_detalhe", "diesel_origem"], dropna=False)
        .agg(valor_total=("valor", "sum"), valor_abs_total=("valor_abs", "sum"), qtd=("valor", "count"))
        .reset_index()
        .sort_values(["grupo_cliente", "placa", "diesel_origem"])
    )

    return detalhe, resumo, resumo_placa


def carregar_media_consumo(cat: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Lê abastecimentos internos (planilha de média) e resume por placa.

    Tratamento: médias muito discrepantes na coluna G são zeradas; a média geral
    por placa usa só valores G > 0 após o tratamento.
    """
    if not XLSX_MEDIA.exists():
        raise FileNotFoundError(f"Planilha de consumo não encontrada: {XLSX_MEDIA}")

    raw = pd.read_excel(XLSX_MEDIA, sheet_name=SHEET_MEDIA, header=None)
    placas_alvo = set(cat["placa_norm"])
    mapa = cat.set_index("placa_norm")[
        ["marca", "grupo_cliente", "contrato_detalhe", "localizacao"]
    ].to_dict("index")

    abastecimentos: list[dict] = []
    resumos_raw: list[dict] = []
    placa_atual: str | None = None

    for _, row in raw.iterrows():
        cel0 = _norm_txt(row[0]).upper()
        is_placa = bool(re.match(r"^EH[LH]\d", cel0))

        if "GERAL" in str(row[5]).upper() and "=" in str(row[5]):
            placa_resumo = cel0 if is_placa else placa_atual
            if placa_resumo and placa_resumo in placas_alvo:
                resumos_raw.append(
                    {"placa": placa_resumo, "media_geral_lh_planilha": parse_numero_br(row[6])}
                )
            continue

        if is_placa and pd.isna(row[1]):
            placa_atual = cel0
            continue

        placa = cel0 if is_placa and pd.notna(row[1]) else placa_atual
        if not placa or placa not in placas_alvo or pd.isna(row[1]):
            continue

        horas = parse_horas_trabalhadas(row[5])
        volume = parse_numero_br(row[2])
        custo = parse_numero_br(row[3])
        media_col_g = parse_numero_br(row[6])  # coluna G — MÉDIA CONSUMO (L/H)
        ref_min, ref_max = parse_numero_br(row[7]), parse_numero_br(row[8])
        info = mapa.get(placa, {})

        consumo_calc = None
        if volume is not None and horas and horas > 0:
            consumo_calc = volume / horas

        media_tratada, media_zerada, motivo_zero = tratar_media_col_g(media_col_g)

        alertas = []
        if media_zerada:
            alertas.append(f"col. G zerada: {motivo_zero}")
        if horimetro_suspeito(row[4]):
            alertas.append("horímetro suspeito (≥9000)")

        abastecimentos.append(
            {
                "placa": placa,
                "marca": info.get("marca", ""),
                "grupo_cliente": info.get("grupo_cliente", ""),
                "contrato_detalhe": info.get("contrato_detalhe", ""),
                "data_abastecimento": pd.to_datetime(row[1], dayfirst=True, errors="coerce"),
                "volume_litros": volume,
                "custo_abastecimento": custo,
                "horimetro": parse_numero_br(row[4]),
                "horas_trabalhadas": horas,
                "media_lh_col_g": media_col_g,
                "media_lh_col_g_tratada": media_tratada,
                "media_col_g_zerada": media_zerada,
                "motivo_zeragem_col_g": motivo_zero,
                "media_lh_calculada": consumo_calc,
                "ref_min_lh": ref_min,
                "ref_max_lh": ref_max,
                "alertas": "; ".join(alertas) if alertas else "",
            }
        )

    det = pd.DataFrame(abastecimentos)
    if det.empty:
        return det, pd.DataFrame(), pd.DataFrame()

    det = det.sort_values(["placa", "data_abastecimento"])

    if resumos_raw:
        res_planilha = (
            pd.DataFrame(resumos_raw)
            .dropna(subset=["placa"])
            .drop_duplicates(subset=["placa"], keep="last")
            .rename(columns={"media_geral_lh_planilha": "media_geral_col_g_planilha_bruta"})
        )
    else:
        res_planilha = pd.DataFrame(columns=["placa", "media_geral_col_g_planilha_bruta"])

    def _agg_placa(g: pd.DataFrame) -> pd.Series:
        g_validos = g[g["media_lh_col_g_tratada"].fillna(0) > 0]
        return pd.Series(
            {
                "qtd_abastecimentos": len(g),
                "qtd_medias_col_g_zeradas": int(g["media_col_g_zerada"].sum()),
                "qtd_medias_col_g_validas": len(g_validos),
                "volume_total_litros": g["volume_litros"].sum(),
                "custo_total_interno": g["custo_abastecimento"].sum(),
                "horas_totais": g["horas_trabalhadas"].sum(),
                "data_primeiro_abast": g["data_abastecimento"].min(),
                "data_ultimo_abast": g["data_abastecimento"].max(),
                "media_geral_lh_tratada": g_validos["media_lh_col_g_tratada"].mean()
                if len(g_validos)
                else None,
                "media_geral_lh_bruta": g["media_lh_col_g"].dropna().mean() if g["media_lh_col_g"].notna().any() else None,
            }
        )

    agg = det.groupby("placa", as_index=False).apply(_agg_placa, include_groups=False).reset_index(drop=True)

    base_placas = cat[["placa_norm", "marca", "grupo_cliente", "contrato_detalhe", "localizacao"]].rename(
        columns={"placa_norm": "placa"}
    )
    res_placa = base_placas.merge(agg, on="placa", how="left").merge(res_planilha, on="placa", how="left")
    res_placa["tem_media_tratada"] = res_placa["qtd_medias_col_g_validas"].fillna(0) > 0

    res_contrato = (
        res_placa.groupby(["grupo_cliente", "contrato_detalhe"], dropna=False)
        .agg(
            qtd_maquinas=("placa", "count"),
            qtd_com_media_tratada=("tem_media_tratada", "sum"),
            medias_col_g_zeradas=("qtd_medias_col_g_zeradas", "sum"),
            volume_litros=("volume_total_litros", "sum"),
            custo_interno_operacional=("custo_total_interno", "sum"),
            media_geral_lh_tratada=("media_geral_lh_tratada", "mean"),
        )
        .reset_index()
        .sort_values(["grupo_cliente", "contrato_detalhe"])
    )

    return det, res_placa, res_contrato


def montar_cruzamento_sagi_media(
    cat: pd.DataFrame,
    detalhe_diesel: pd.DataFrame,
    res_placa_media: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Cruza diesel interno SAGI com abastecimento interno da planilha de médias."""
    sagi_placa = (
        detalhe_diesel.groupby(["placa", "marca", "grupo_cliente", "contrato_detalhe", "diesel_origem"], dropna=False)
        .agg(valor_sagi=("valor", "sum"), valor_abs_sagi=("valor_abs", "sum"), qtd_lanc_sagi=("valor", "count"))
        .reset_index()
    )
    sagi_pivot = (
        sagi_placa.pivot_table(
            index=["placa", "marca", "grupo_cliente", "contrato_detalhe"],
            columns="diesel_origem",
            values="valor_abs_sagi",
            aggfunc="sum",
            fill_value=0,
        )
        .reset_index()
    )
    for col in ("Externo", "Interno"):
        if col not in sagi_pivot.columns:
            sagi_pivot[col] = 0.0
    sagi_pivot["custo_sagi_total"] = sagi_pivot["Externo"] + sagi_pivot["Interno"]

    media = res_placa_media[
        [
            "placa",
            "qtd_abastecimentos",
            "qtd_medias_col_g_zeradas",
            "qtd_medias_col_g_validas",
            "volume_total_litros",
            "custo_total_interno",
            "media_geral_lh_tratada",
            "media_geral_lh_bruta",
            "media_geral_col_g_planilha_bruta",
            "data_primeiro_abast",
            "data_ultimo_abast",
        ]
    ].copy()

    # Catálogo como espinha: garante marca/contrato mesmo sem lançamento SAGI no período
    base_placas = cat[
        ["placa_norm", "marca", "grupo_cliente", "contrato_detalhe", "localizacao"]
    ].rename(columns={"placa_norm": "placa"})
    sagi_cols = [c for c in sagi_pivot.columns if c not in {"marca", "grupo_cliente", "contrato_detalhe"}]
    cruz = (
        base_placas.merge(sagi_pivot[sagi_cols], on="placa", how="left")
        .merge(media, on="placa", how="left")
    )

    cruz["custo_sagi_interno_periodo_operacional"] = 0.0
    if not detalhe_diesel.empty and detalhe_diesel["data_nf"].notna().any():
        det = detalhe_diesel.copy()
        det["data_nf"] = pd.to_datetime(det["data_nf"], errors="coerce")
        det_interno = det[det["diesel_origem"] == "Interno"]
        for _, row in cruz.iterrows():
            if pd.isna(row.get("data_primeiro_abast")) or pd.isna(row.get("data_ultimo_abast")):
                continue
            mask = (
                (det_interno["placa"] == row["placa"])
                & (det_interno["data_nf"] >= row["data_primeiro_abast"])
                & (det_interno["data_nf"] <= row["data_ultimo_abast"])
            )
            cruz.loc[cruz["placa"] == row["placa"], "custo_sagi_interno_periodo_operacional"] = float(
                det_interno.loc[mask, "valor_abs"].sum()
            )
    cruz["diff_sagi_interno_vs_operacional"] = cruz["Interno"] - cruz["custo_total_interno"].fillna(0)
    cruz["diff_sagi_interno_periodo_vs_operacional"] = (
        cruz["custo_sagi_interno_periodo_operacional"] - cruz["custo_total_interno"].fillna(0)
    )
    cruz["pct_cobertura_interno"] = cruz.apply(
        lambda r: r["custo_total_interno"] / r["custo_sagi_interno_periodo_operacional"]
        if r.get("custo_sagi_interno_periodo_operacional") and r["custo_sagi_interno_periodo_operacional"] > 0
        and pd.notna(r.get("custo_total_interno"))
        else None,
        axis=1,
    )
    cruz["sem_dados_operacionais"] = cruz["qtd_abastecimentos"].isna() | (cruz["qtd_abastecimentos"] == 0)
    cruz = cruz.sort_values(["grupo_cliente", "placa"])

    ordem = [
        "placa",
        "marca",
        "grupo_cliente",
        "contrato_detalhe",
        "localizacao",
        "qtd_abastecimentos",
        "volume_total_litros",
        "custo_total_interno",
        "media_geral_lh_tratada",
        "media_geral_lh_bruta",
        "qtd_medias_col_g_zeradas",
        "qtd_medias_col_g_validas",
        "data_primeiro_abast",
        "data_ultimo_abast",
        "Interno",
        "Externo",
        "custo_sagi_total",
        "custo_sagi_interno_periodo_operacional",
        "diff_sagi_interno_vs_operacional",
        "diff_sagi_interno_periodo_vs_operacional",
        "pct_cobertura_interno",
        "media_geral_col_g_planilha_bruta",
        "sem_dados_operacionais",
    ]
    cruz = cruz[[c for c in ordem if c in cruz.columns]]

    res_contrato = (
        cruz.groupby(["grupo_cliente", "contrato_detalhe"], dropna=False)
        .agg(
            qtd_maquinas=("placa", "count"),
            custo_sagi_total=("custo_sagi_total", "sum"),
            custo_sagi_externo=("Externo", "sum"),
            custo_sagi_interno=("Interno", "sum"),
            custo_interno_operacional=("custo_total_interno", "sum"),
            volume_litros_operacional=("volume_total_litros", "sum"),
            media_geral_lh_tratada=("media_geral_lh_tratada", "mean"),
            medias_col_g_zeradas=("qtd_medias_col_g_zeradas", "sum"),
            qtd_sem_dados_operacionais=("sem_dados_operacionais", "sum"),
        )
        .reset_index()
        .sort_values(["grupo_cliente", "contrato_detalhe"])
    )

    return cruz, res_contrato


def _estilizar_aba(ws, col_formats: dict[int, str] | None = None, header_row: int = 1):
    col_formats = col_formats or {}
    for cell in ws[header_row]:
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row):
        for cell in row:
            fmt = col_formats.get(cell.column)
            if fmt:
                cell.number_format = fmt
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        max_len = 0
        for cell in ws[letter]:
            if cell.value is not None:
                max_len = max(max_len, min(len(str(cell.value)), 55))
        ws.column_dimensions[letter].width = max(10, max_len + 2)
    ws.freeze_panes = f"A{header_row + 1}"


def _rotulo_filtros() -> str:
    return (
        f"Recorte: {PERIODO_INI.strftime('%m/%Y')} a {PERIODO_FIM.strftime('%m/%Y')} | "
        f"Escavadeiras Hyundai/Liebherr ({len(PARQUE_ESCAVADEIRAS)} placas no parque oficial)"
    )


def _escrever_aba_com_rotulo(
    writer,
    sheet_name: str,
    rotulo: str,
    tabelas: list[tuple[str, pd.DataFrame]],
):
    """Escreve uma aba com rótulo na linha 1 e várias tabelas empilhadas."""
    start = 2
    ws = None
    for titulo, df in tabelas:
        if df is None or df.empty:
            continue
        if ws is None:
            df.to_excel(writer, sheet_name=sheet_name, index=False, startrow=start)
            ws = writer.sheets[sheet_name]
            ws.cell(row=1, column=1, value=rotulo)
            ws.cell(row=1, column=1).font = Font(bold=True, italic=True, size=10)
            start += len(df) + 3
        else:
            ws.cell(row=start, column=1, value=titulo)
            ws.cell(row=start, column=1).font = Font(bold=True, size=12)
            df.to_excel(writer, sheet_name=sheet_name, index=False, startrow=start)
            start += len(df) + 3


def exportar_excel(
    cat: pd.DataFrame,
    resumo_cat: pd.DataFrame,
    detalhe: pd.DataFrame,
    resumo: pd.DataFrame,
    resumo_placa: pd.DataFrame,
    out_path: Path = XLSX_OUT,
):
    rotulo = _rotulo_filtros()
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        start_cat = 2
        cat.to_excel(writer, sheet_name="Maquinas_por_Contrato", index=False, startrow=start_cat)
        start_resumo = start_cat + len(cat) + 3
        ws1 = writer.sheets["Maquinas_por_Contrato"]
        ws1.cell(row=1, column=1, value=rotulo)
        ws1.cell(row=1, column=1).font = Font(bold=True, italic=True, size=10)
        ws1.cell(row=start_resumo, column=1, value="RESUMO POR CONTRATO")
        ws1.cell(row=start_resumo, column=1).font = Font(bold=True, size=12)
        resumo_cat.to_excel(writer, sheet_name="Maquinas_por_Contrato", index=False, startrow=start_resumo)

        start_diesel = 2
        resumo.to_excel(writer, sheet_name="Diesel_Interno_Externo", index=False, startrow=start_diesel)
        start_d = start_diesel + len(resumo) + 3
        ws2 = writer.sheets["Diesel_Interno_Externo"]
        ws2.cell(row=1, column=1, value=rotulo)
        ws2.cell(row=1, column=1).font = Font(bold=True, italic=True, size=10)
        ws2.cell(row=start_d, column=1, value="RESUMO POR PLACA")
        ws2.cell(row=start_d, column=1).font = Font(bold=True, size=12)
        resumo_placa.to_excel(writer, sheet_name="Diesel_Interno_Externo", index=False, startrow=start_d)
        start_det = start_d + len(resumo_placa) + 3
        ws2.cell(row=start_det, column=1, value="DETALHE DOS LANÇAMENTOS")
        ws2.cell(row=start_det, column=1).font = Font(bold=True, size=12)
        detalhe.to_excel(writer, sheet_name="Diesel_Interno_Externo", index=False, startrow=start_det)

    from openpyxl import load_workbook

    wb = load_workbook(out_path)
    ws1 = wb["Maquinas_por_Contrato"]
    brl_cols_cat = {cat.columns.get_loc("depreciacao_mensal") + 1 + 0: FMT_BRL}
    _estilizar_aba(ws1, brl_cols_cat, header_row=3)

    for sheet in ("Diesel_Interno_Externo",):
        if sheet in wb.sheetnames:
            _estilizar_aba(wb[sheet], header_row=3)

    wb.save(out_path)


def main():
    print(_rotulo_filtros())
    print("Carregando catálogo de máquinas...")
    cat = carregar_catalogo_maquinas()
    resumo_cat = montar_resumo_contratos(cat)
    print(f"  Escavadeiras H/L: {len(cat)} | Grupos: {cat['grupo_cliente'].nunique()}")
    print(cat.groupby("marca")["placa"].count().to_string())

    print("Carregando diesel da base de fechamento...")
    detalhe, resumo, resumo_placa = carregar_diesel(cat)
    print(f"  Lançamentos diesel (período + H/L): {len(detalhe)}")
    print(resumo.groupby("diesel_origem")["valor_abs_total"].sum().to_string())

    candidatos = [XLSX_OUT, CUSTOS / "analise_contratos_diesel_maquinas_novo.xlsx"]
    for out_path in candidatos:
        try:
            print(f"Exportando {out_path}...")
            exportar_excel(
                cat,
                resumo_cat,
                detalhe,
                resumo,
                resumo_placa,
                out_path=out_path,
            )
            print("Concluído.")
            break
        except PermissionError:
            if out_path == candidatos[-1]:
                raise
            print(f"  {out_path.name} aberto no Excel — tentando arquivo alternativo...")


if __name__ == "__main__":
    main()
