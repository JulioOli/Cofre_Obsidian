"""Auditoria completa: dashboard vs base consolidada."""
from __future__ import annotations

import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "outputs" / "relatorios" / "comparativo_liebherr_hyundai"))

from openpyxl import load_workbook

from atualizar_abas_gasto_filtrado import (
    LINHAS_GASTO,
    carregar_base,
    carregar_diesel_2026_q1_media,
    montar_lookup,
    CONTAS_GERAL,
    CONTAS_MANUTENCAO,
    CONTAS_COMBUSTIVEL,
    PARQUE,
    ALL_MAQUINAS,
    soma_maquinas,
)
from export_dashboard_excel import (
    MAQ_ID_RE,
    LOCAL_EXCEL_TO_DF,
    carregar_df_maq,
    montar_kpis,
    montar_dados_mix_contas_agrupado,
    montar_fonte_dados,
    media_mensal_grupo,
    _maquinas_local_marca,
    _gasto_maquina_ano,
)

OUT = ROOT / "outputs" / "tabelas" / "dashboard_liebherr_hyundai.xlsx"
TOL = 0.05

issues: list[str] = []


def chk(cond: bool, msg: str) -> None:
    if not cond:
        issues.append(msg)


def main() -> int:
    df_base = carregar_base()
    diesel = carregar_diesel_2026_q1_media(df_base)
    lookups = {
        "geral": montar_lookup(df_base, CONTAS_GERAL, diesel, "add"),
        "manut": montar_lookup(df_base, CONTAS_MANUTENCAO),
        "comb": montar_lookup(df_base, CONTAS_COMBUSTIVEL, diesel, "combustivel"),
    }

    wb = load_workbook(OUT, data_only=True)
    wb_f = load_workbook(OUT, data_only=False)

    for i, (local, marca, _q, maqs) in enumerate(LINHAS_GASTO):
        r = 3 + i
        exp25 = round(soma_maquinas(maqs, lookups["geral"], 2025), 2)
        exp26 = round(soma_maquinas(maqs, lookups["geral"], 2026), 2)
        g25 = wb["Gasto geral"].cell(r, 4).value or 0
        g26 = wb["Gasto geral"].cell(r, 7).value or 0
        chk(abs(g25 - exp25) <= TOL, f"Gasto geral [{local}/{marca}] 2025: {g25} vs {exp25}")
        chk(abs(g26 - exp26) <= TOL, f"Gasto geral [{local}/{marca}] 2026: {g26} vs {exp26}")

    wsf = wb_f["Gasto geral"]
    for i in range(len(LINHAS_GASTO)):
        r = 3 + i
        chk(wsf.cell(r, 6).value == f"=D{r}/12", f"Gasto geral r{r} média 2025 incorreta")
        chk(wsf.cell(r, 9).value == f"=G{r}/3", f"Gasto geral r{r} média 2026 incorreta")

    for r in range(1, 55):
        if wsf.cell(r, 2).value == "HYUNDAI" and isinstance(wsf.cell(r, 4).value, str):
            f6 = wsf.cell(r, 6).value or ""
            chk("/12" not in str(f6), f"Resumo HYUNDAI r{r}: média 2025 divide por 12 em excesso ({f6})")
            chk("D10" in str(wsf.cell(r, 4).value), f"Resumo HYUNDAI r{r}: falta Maringá Cidade Alta (D10)")
            break

    for sheet, lk, start in [
        ("Gastos com Manutenção", "manut", 5),
        ("Gastos com combustível", "comb", 5),
    ]:
        ws = wb[sheet]
        for i, (local, marca, _q, maqs) in enumerate(LINHAS_GASTO):
            r = start + i
            exp25 = round(soma_maquinas(maqs, lookups[lk], 2025), 2)
            exp26 = round(soma_maquinas(maqs, lookups[lk], 2026), 2)
            chk(abs((ws.cell(r, 4).value or 0) - exp25) <= TOL, f"{sheet} [{local}] 2025")
            chk(abs((ws.cell(r, 7).value or 0) - exp26) <= TOL, f"{sheet} [{local}] 2026")

    manut = sum(lookups["manut"].get((m, a), 0) for m in ALL_MAQUINAS for a in (2025, 2026))
    comb = sum(lookups["comb"].get((m, a), 0) for m in ALL_MAQUINAS for a in (2025, 2026))
    geral = sum(lookups["geral"].get((m, a), 0) for m in ALL_MAQUINAS for a in (2025, 2026))
    chk(abs(manut + comb - geral) <= 1, f"Manut+Comb={manut+comb:.0f} vs Geral={geral:.0f}")

    df_maq = carregar_df_maq()
    dados_mix = montar_dados_mix_contas_agrupado(df_maq)
    ws = wb["Mix por conta"]
    for r in range(7, 10):
        label = ws.cell(r, 2).value
        if label in dados_mix:
            d = dados_mix[label]
            chk(abs((ws.cell(r, 3).value or 0) - d["rpm_LIEBHERR"]) < 0.2, f"Mix {label} LIEBHERR rpm")
            chk(abs((ws.cell(r, 15).value or 0) - d["pct_LIEBHERR"]) < 0.2, f"Mix {label} LIEBHERR %")

    if "Gasto por Local - Detalhado" in wb.sheetnames:
        ws = wb["Gasto por Local - Detalhado"]
        for r in range(4, ws.max_row + 1):
            mid = ws.cell(r, 3).value
            if not isinstance(mid, str) or not MAQ_ID_RE.match(mid.strip()):
                continue
            mid = mid.strip().upper()
            for ano, col in ((2025, 4), (2026, 5)):
                exp = round(_gasto_maquina_ano(df_maq, mid, ano), 2)
                got = ws.cell(r, col).value or 0
                chk(abs(got - exp) <= TOL, f"Detalhado {mid} {ano}: {got} vs {exp}")

    if "Resumo Local×Marca" in wb.sheetnames:
        agg = media_mensal_grupo(df_maq, ["local_operacao", "MARCA"])
        ws = wb["Resumo Local×Marca"]
        for r in range(4, ws.max_row + 1):
            loc_ex = ws.cell(r, 1).value
            marca = ws.cell(r, 2).value
            if not loc_ex or not marca:
                continue
            local = LOCAL_EXCEL_TO_DF.get(str(loc_ex).strip(), str(loc_ex).strip())
            for ano, col in ((2025, 5), (2026, 7)):
                sub = agg[(agg["local_operacao"] == local) & (agg["MARCA"] == marca) & (agg["ano_nf"] == ano)]
                exp = round(float(sub["gasto_total"].iloc[0]), 2) if len(sub) else 0.0
                got = ws.cell(r, col).value or 0
                chk(abs(got - exp) <= 1, f"Resumo {loc_ex}/{marca} {ano}: {got} vs {exp}")

    kpis = dict(montar_kpis(df_maq))
    ws = wb["Visão Geral"]
    for r in range(1, ws.max_row + 1):
        lb = ws.cell(r, 1).value
        if lb not in ("LIEBHERR", "HYUNDAI", "Comparativo justo"):
            continue
        ev, v = ws.cell(r, 2).value, kpis.get(lb, "")
        if not ev or not v:
            continue
        nums_ev = re.findall(r"[\d.]+", str(ev).replace(",", "."))
        nums_v = re.findall(r"[\d.]+", str(v).replace(",", "."))
        if nums_ev and nums_v:
            chk(abs(float(nums_ev[0]) - float(nums_v[0])) < 2, f"KPI {lb}: {ev} vs {v}")

    if "Fonte_Dados" in wb.sheetnames:
        exp = montar_fonte_dados(df_maq)["Gasto_Abs"].sum()
        ws = wb["Fonte_Dados"]
        col = next((c for c in range(1, 30) if ws.cell(3, c).value == "Gasto_Abs"), None)
        if col:
            got = sum(ws.cell(r, col).value or 0 for r in range(4, ws.max_row + 1))
            chk(abs(got - exp) < 1, f"Fonte_Dados soma: {got:.0f} vs {exp:.0f}")

    print(f"Auditoria: {len(issues)} problema(s)")
    for msg in issues:
        print(" -", msg)
    return len(issues)


if __name__ == "__main__":
    raise SystemExit(main())
