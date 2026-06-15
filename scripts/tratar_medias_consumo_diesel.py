"""
Gera planilha de médias coerentes somente a partir da planilha operacional de diesel interno.

Fonte:
  - 02-Referencias/Custos-Maquinas/Media_de_Consumo_de_Diesel_das_Maquinas.xlsx

Saída:
  - 02-Referencias/Custos-Maquinas/medias_consumo_diesel_tratado.xlsx

Não cruza com SAGI, catálogo de contratos nem base de fechamento.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from medias_consumo_diesel import (
    LH_MAX_COL_G,
    LH_MIN_COL_G,
    XLSX_MEDIA,
    carregar_planilha_media,
)

ROOT = Path(__file__).resolve().parents[1]
CUSTOS = ROOT / "02-Referencias" / "Custos-Maquinas"
XLSX_OUT = CUSTOS / "medias_consumo_diesel_tratado.xlsx"

FONT_HEADER = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
FILL_HEADER = PatternFill("solid", fgColor="305496")


def _rotulo() -> str:
    return (
        f"Abastecimento interno — fonte: {XLSX_MEDIA.name} | "
        f"Tratamento: col. G zerada se < {LH_MIN_COL_G} ou > {LH_MAX_COL_G} L/h | "
        f"Média geral = média aritmética dos G tratados (> 0) por placa"
    )


def _estilizar_aba(ws, header_rows: list[int]):
    for header_row in header_rows:
        for cell in ws[header_row]:
            cell.font = FONT_HEADER
            cell.fill = FILL_HEADER
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        max_len = max(
            (min(len(str(c.value or "")), 50) for c in ws[letter]),
            default=10,
        )
        ws.column_dimensions[letter].width = max(10, max_len + 2)
    if header_rows:
        ws.freeze_panes = f"A{header_rows[0] + 1}"


def exportar(det: pd.DataFrame, res_placa: pd.DataFrame, resumo: pd.DataFrame, out_path: Path):
    media_zeradas = det[det["media_col_g_zerada"]].copy()
    media_validas = det[~det["media_col_g_zerada"]].copy()
    header_rows = [3]

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        sheet = "Medias_Tratadas"
        start = 2
        resumo.to_excel(writer, sheet_name=sheet, index=False, startrow=start)
        ws = writer.sheets[sheet]
        ws.cell(row=1, column=1, value=_rotulo())
        ws.cell(row=1, column=1).font = Font(bold=True, italic=True, size=10)

        secoes = [
            ("RESUMO POR PLACA", res_placa),
            ("DETALHE ABASTECIMENTOS", det),
            ("ABASTECIMENTOS COM MÉDIA VÁLIDA (COL. G)", media_validas),
            ("MÉDIAS COL. G ZERADAS", media_zeradas),
        ]
        start = start + len(resumo) + 3
        for titulo, df in secoes:
            if df.empty:
                continue
            ws.cell(row=start, column=1, value=titulo)
            ws.cell(row=start, column=1).font = Font(bold=True, size=12)
            df.to_excel(writer, sheet_name=sheet, index=False, startrow=start)
            header_rows.append(start + 1)
            start += len(df) + 3

    wb = load_workbook(out_path)
    _estilizar_aba(wb["Medias_Tratadas"], header_rows)
    wb.save(out_path)


def main():
    print(_rotulo())
    det, res_placa, resumo = carregar_planilha_media()
    print(f"  Abastecimentos: {len(det)} | Placas: {res_placa['placa'].nunique()}")
    print(f"  Médias col. G zeradas: {int(det['media_col_g_zerada'].sum())}")
    print(
        f"  Placas com média tratada: "
        f"{int((res_placa['qtd_medias_col_g_validas'].fillna(0) > 0).sum())}"
    )

    candidatos = [XLSX_OUT, CUSTOS / "medias_consumo_diesel_tratado_novo.xlsx"]
    for out_path in candidatos:
        try:
            print(f"Exportando {out_path}...")
            exportar(det, res_placa, resumo, out_path)
            print("Concluído.")
            break
        except PermissionError:
            if out_path == candidatos[-1]:
                raise
            print(f"  {out_path.name} aberto — tentando arquivo alternativo...")


if __name__ == "__main__":
    main()
